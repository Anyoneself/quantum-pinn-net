#!/usr/bin/env python3
"""Build a journal-style editable DOCX from the submission Markdown.

Preset: compact_reference_guide
Named overrides: gxnu_journal_article (A4, single column, black-and-white,
Chinese/English captions, three-line tables, English extended abstract at end).
"""

from __future__ import annotations

import re
import os
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[3]
SOURCE = ROOT / "docs" / "QGE_PINN_GXNU_STYLE_SUBMISSION.md"
OUTPUT = ROOT / "v2/paper/output/docx/QGE_PINN_投稿稿_广西师大学报风格.docx"

# Arial Unicode MS is available to the bundled LibreOffice renderer and covers
# simplified Chinese reliably; Word users can replace it with a journal font.
CN_BODY = "Arial Unicode MS"
CN_BOLD = "Arial Unicode MS"
EN_BODY = "Times New Roman"
MATH_FONT = "Cambria Math"


def set_run_font(run, *, size=10.5, bold=False, cn=CN_BODY, en=EN_BODY, italic=False):
    # LibreOffice on macOS may ignore w:eastAsia when w:ascii points to a font
    # without CJK glyphs. Use the CJK-capable face for mixed Chinese body runs;
    # English-only sections pass Times New Roman as ``cn`` explicitly.
    effective = cn
    run.font.name = effective
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cn)
    run._element.rPr.rFonts.set(qn("w:ascii"), effective)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), effective)
    return run


def set_cell_margins(cell, top=45, start=80, bottom=45, end=80):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **edges):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge_name, attrs in edges.items():
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        for key, value in attrs.items():
            edge.set(qn(f"w:{key}"), str(value))


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trPr.append(tbl_header)


def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    trPr.append(cant_split)


def set_keep_with_next(paragraph, value=True):
    pPr = paragraph._p.get_or_add_pPr()
    keep = pPr.find(qn("w:keepNext"))
    if keep is None:
        keep = OxmlElement("w:keepNext")
        pPr.append(keep)
    keep.set(qn("w:val"), "1" if value else "0")


def set_widow_control(paragraph, value=True):
    pPr = paragraph._p.get_or_add_pPr()
    node = pPr.find(qn("w:widowControl"))
    if node is None:
        node = OxmlElement("w:widowControl")
        pPr.append(node)
    node.set(qn("w:val"), "1" if value else "0")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, size=8.5)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.65)
    section.bottom_margin = Cm(1.65)
    section.left_margin = Cm(2.05)
    section.right_margin = Cm(2.05)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = EN_BODY
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CN_BODY)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(16)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Pt(21)

    styles = doc.styles
    for name, level, size in (
        ("Journal Heading 1", 1, 13),
        ("Journal Heading 2", 2, 10.5),
        ("Journal Heading 3", 3, 10.5),
    ):
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        style.font.name = EN_BODY
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CN_BOLD)
        style.font.size = Pt(size)
        style.font.bold = level >= 2
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_before = Pt(8 if level == 1 else 5)
        style.paragraph_format.space_after = Pt(3 if level == 1 else 1)
        style.paragraph_format.keep_with_next = True

    caption = styles.add_style("Journal Caption", WD_STYLE_TYPE.PARAGRAPH)
    caption.base_style = normal
    caption.font.name = EN_BODY
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), CN_BODY)
    caption.font.size = Pt(8.5)
    caption.paragraph_format.first_line_indent = Pt(0)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    caption.paragraph_format.space_before = Pt(1)
    caption.paragraph_format.space_after = Pt(1)
    caption.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hp.paragraph_format.space_after = Pt(0)
    set_run_font(hp.add_run("投稿稿　量子门等价特征增强物理引导神经网络"), size=8.5)
    add_page_number(section.footer.paragraphs[0])


def clean_inline(text: str) -> str:
    replacements = {
        r"\widetilde": "",
        r"\rightarrow": "→",
        r"\mathbb": "",
        r"\in": "∈",
        r"\nabla": "∇",
        r"\mu": "μ",
        r"\phi": "φ",
        r"\theta": "θ",
        r"\eta": "η",
        r"\lambda": "λ",
        r"\varepsilon": "ε",
        r"\sum": "Σ",
        r"\sqrt": "√",
        r"\operatorname": "",
        r"\left": "",
        r"\right": "",
        r"\qquad": "　",
        r"\,": "",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = text.replace("$", "").replace("\\", "")
    return text


def add_inline(paragraph, text: str, *, size=10.5, base_bold=False, cn=CN_BODY, en=EN_BODY):
    text = clean_inline(text)
    pieces = re.split(r"(\*\*.*?\*\*)", text)
    for piece in pieces:
        if not piece:
            continue
        bold = base_bold
        if piece.startswith("**") and piece.endswith("**"):
            piece = piece[2:-2]
            bold = True
        set_run_font(paragraph.add_run(piece), size=size, bold=bold, cn=cn, en=en)


def add_body_paragraph(doc, text, *, indent=True, size=10.5, bold_prefix=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(16)
    p.paragraph_format.first_line_indent = Pt(21 if indent else 0)
    p.paragraph_format.space_after = Pt(0)
    set_widow_control(p)
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(
            p.add_run(clean_inline(bold_prefix.replace("**", ""))),
            size=size,
            bold=True,
            cn=CN_BOLD,
        )
        add_inline(p, text[len(bold_prefix):], size=size)
    else:
        add_inline(p, text, size=size)
    return p


def add_equation(doc, text):
    match = re.search(r"\\tag\{(\d+)\}", text)
    tag = match.group(1) if match else ""
    latex = re.sub(r"\\tag\{\d+\}", "", text).strip()
    equation_dir = ROOT / "tmp" / "equation_images"
    equation_dir.mkdir(parents=True, exist_ok=True)
    key = f"eq_{tag or 'x'}_{abs(hash(latex))}"
    image_path = equation_dir / f"{key}.png"
    if not image_path.exists():
        os.environ.setdefault("MPLCONFIGDIR", "/private/tmp/codex-mpl-cache")
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt

        # Wider canvas for long loss/IDW expressions; tight cropping removes
        # unused space while retaining the equation number at the right edge.
        width = 11.0 if len(latex) > 120 else (8.2 if len(latex) > 65 else 6.4)
        fig = plt.figure(figsize=(width, 0.48), dpi=220, facecolor="white")
        fig.text(0.48, 0.5, f"${latex}$", ha="center", va="center", fontsize=10.5)
        if tag:
            fig.text(0.985, 0.5, f"({tag})", ha="right", va="center", fontsize=10.5)
        fig.savefig(image_path, dpi=220, bbox_inches="tight", pad_inches=0.03, facecolor="white")
        plt.close(fig)
    with Image.open(image_path) as im:
        w_px, h_px = im.size
    width_cm = min(15.5, max(5.5, w_px / 220 * 2.54))
    height_cm = width_cm * h_px / w_px
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.keep_together = True
    p.add_run().add_picture(str(image_path), width=Cm(width_cm), height=Cm(height_cm))
    return p


def add_figure(doc, image_path: Path):
    with Image.open(image_path) as im:
        w, h = im.size
    max_w_cm = 15.8
    max_h_cm = 8.3 if image_path.name == "method_flowchart.png" else 7.5
    ratio = w / h
    width_cm = min(max_w_cm, max_h_cm * ratio)
    height_cm = width_cm / ratio
    if height_cm > max_h_cm:
        height_cm = max_h_cm
        width_cm = height_cm * ratio
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.keep_together = True
    set_keep_with_next(p)
    p.add_run().add_picture(str(image_path), width=Cm(width_cm), height=Cm(height_cm))
    return p


def parse_table(lines, start):
    rows = []
    idx = start
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        cells = [c.strip() for c in lines[idx].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
            rows.append(cells)
        idx += 1
    return rows, idx


def add_table(doc, rows):
    cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.autofit = False
    total_width_cm = 16.2
    first_width = 4.0 if cols >= 5 else total_width_cm / cols
    other_width = (total_width_cm - first_width) / (cols - 1) if cols > 1 else total_width_cm
    widths = [first_width] + [other_width] * (cols - 1)
    for r_idx, row in enumerate(rows):
        prevent_row_split(table.rows[r_idx])
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.width = Cm(widths[c_idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            add_inline(p, value, size=8.2, base_bold=(r_idx == 0), cn=CN_BOLD if r_idx == 0 else CN_BODY)
            if r_idx < len(rows) - 1:
                set_keep_with_next(p)
            none = {"val": "nil"}
            set_cell_border(cell, top=none, bottom=none, start=none, end=none, insideH=none, insideV=none)
            if r_idx == 0:
                set_cell_border(
                    cell,
                    top={"val": "single", "sz": "12", "color": "000000"},
                    bottom={"val": "single", "sz": "6", "color": "000000"},
                )
            elif r_idx == len(rows) - 1:
                set_cell_border(cell, bottom={"val": "single", "sz": "12", "color": "000000"})
    set_repeat_table_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Journal Heading {min(level, 3)}")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_inline(p, text, size=13 if level == 1 else 10.5, base_bold=(level >= 2), cn=CN_BOLD)
    set_keep_with_next(p)
    return p


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)

    # First-page journal masthead and metadata placeholders.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    set_run_font(p.add_run("广西师范大学学报（自然科学版）投稿稿"), size=10.5, bold=True, cn=CN_BOLD)

    title = lines[0].lstrip("# ").strip()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(7)
    set_run_font(p.add_run(title), size=18, bold=True, cn=CN_BOLD)

    for text, size, bold in ((lines[2], 11, False), (lines[4], 9.5, False)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        add_inline(p, text, size=size, base_bold=bold)

    idx = 6
    while idx < len(lines) and not lines[idx].startswith("## 摘要"):
        idx += 1
    idx += 1
    while idx < len(lines):
        raw = lines[idx].strip()
        if raw.startswith("## English title"):
            doc.add_page_break()
            add_english_abstract(doc, lines, idx + 1)
            break
        if not raw:
            idx += 1
            continue
        if raw.startswith("!["):
            rel = re.search(r"\((.+)\)", raw).group(1)
            add_figure(doc, (SOURCE.parent / rel).resolve())
            idx += 1
            continue
        if raw.startswith("|"):
            rows, idx = parse_table(lines, idx)
            add_table(doc, rows)
            continue
        if raw == "$$":
            eq_lines = []
            idx += 1
            while idx < len(lines) and lines[idx].strip() != "$$":
                eq_lines.append(lines[idx].strip())
                idx += 1
            add_equation(doc, " ".join(eq_lines))
            idx += 1
            continue
        if raw.startswith("### "):
            add_heading(doc, raw[4:], 2)
            idx += 1
            continue
        if raw.startswith("## "):
            heading = raw[3:]
            if heading == "摘要":
                idx += 1
                continue
            if heading == "参考文献":
                add_heading(doc, "参 考 文 献", 1)
            else:
                add_heading(doc, heading, 1)
            idx += 1
            continue
        if raw.startswith("**图") or raw.startswith("**Fig.") or raw.startswith("**表") or raw.startswith("**Tab."):
            p = doc.add_paragraph(style="Journal Caption")
            add_inline(p, raw, size=8.5, base_bold=raw.startswith(("**表", "**Tab.")))
            # Keep Chinese captions with the bilingual line; keep table captions
            # with the following table. The final English figure caption may
            # break before the next body paragraph.
            set_keep_with_next(p, not raw.startswith("**Fig."))
            idx += 1
            continue
        if raw.startswith("**关键词："):
            add_body_paragraph(doc, raw, indent=False, size=9, bold_prefix="**关键词：**")
            idx += 1
            continue
        if raw.startswith("**中图分类号："):
            add_body_paragraph(doc, raw, indent=False, size=9)
            meta = doc.add_paragraph()
            meta.alignment = WD_ALIGN_PARAGRAPH.LEFT
            meta.paragraph_format.first_line_indent = Pt(0)
            meta.paragraph_format.space_before = Pt(4)
            meta.paragraph_format.space_after = Pt(5)
            add_inline(meta, "收稿日期：____-__-__　修回日期：____-__-__　　基金项目：待补充　　*通信作者：待补充", size=8)
            idx += 1
            continue
        if raw.startswith("[") and re.match(r"\[\d+\]", raw):
            p = add_body_paragraph(doc, raw, indent=False, size=8.5)
            p.paragraph_format.left_indent = Pt(18)
            p.paragraph_format.first_line_indent = Pt(-18)
            p.paragraph_format.line_spacing = Pt(12)
            continue_idx = idx + 1
            idx = continue_idx
            continue
        if raw.startswith("**") and "：" in raw:
            prefix = raw.split("**", 2)[1] + "**"
            add_body_paragraph(doc, raw, indent=False, size=9, bold_prefix="**" + prefix)
            idx += 1
            continue

        # Abstract is the first ordinary paragraph after the heading.
        is_abstract = idx > 6 and not any(p.text.startswith("摘要：") for p in doc.paragraphs)
        if is_abstract:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            p.paragraph_format.line_spacing = Pt(14)
            set_run_font(p.add_run("摘要："), size=9, bold=True, cn=CN_BOLD)
            add_inline(p, raw, size=9)
        else:
            add_body_paragraph(doc, raw)
        idx += 1

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = title
    doc.core_properties.subject = "非均质渗透率场空间补全；独立测试；五随机种子消融研究"
    doc.core_properties.author = "作者信息待补充"
    doc.core_properties.keywords = "渗透率场；物理引导神经网络；量子启发特征；残差自适应"
    doc.save(OUTPUT)
    print(OUTPUT)


def add_english_abstract(doc, lines, start):
    idx = start
    while idx < len(lines) and not lines[idx].strip():
        idx += 1
    title = lines[idx].lstrip("# ").strip()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(5)
    set_run_font(p.add_run(title), size=15, bold=True, cn=EN_BODY)
    idx += 1
    for _ in range(2):
        while idx < len(lines) and not lines[idx].strip():
            idx += 1
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        add_inline(p, lines[idx].strip(), size=10 if _ == 0 else 9)
        idx += 1
    while idx < len(lines):
        raw = lines[idx].strip()
        if raw:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            p.paragraph_format.line_spacing = Pt(13)
            if raw.startswith("**Abstract:**"):
                set_run_font(p.add_run("Abstract: "), size=9, bold=True, cn=EN_BODY)
                add_inline(p, raw[len("**Abstract:**"):].strip(), size=9, cn=EN_BODY)
            elif raw.startswith("**Keywords:**"):
                set_run_font(p.add_run("Keywords: "), size=9, bold=True, cn=EN_BODY)
                add_inline(p, raw[len("**Keywords:**"):].strip(), size=9, cn=EN_BODY)
            else:
                add_inline(p, raw, size=9, cn=EN_BODY)
        idx += 1


if __name__ == "__main__":
    build()
