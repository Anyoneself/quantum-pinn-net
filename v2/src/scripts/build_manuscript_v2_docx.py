from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
import latex2mathml.converter
import mathml2omml

import build_submission_docx as common
import build_chinese_core_docx as style

ROOT = Path(__file__).resolve().parents[3]
ASSETS = {}


def text_run(p, text, size=10.5, bold=False):
    font = "Songti SC" if re.search(r"[\u3400-\u9fff]", text) else "Times New Roman"
    return common.set_run_font(p.add_run(text), size=size, bold=bold, cn=font)


def inline(p, text, size=10.5):
    for piece in re.split(r"(\\\(.*?\\\)|\*\*.*?\*\*)", text):
        if piece.startswith(r"\(") and piece.endswith(r"\)"):
            latex = piece[2:-2]
            xml = mathml2omml.convert(latex2mathml.converter.convert(latex))
            try:
                wrapper = etree.fromstring((
                    '<root xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
                    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                    + xml + "</root>"
                ).encode())
            except etree.XMLSyntaxError:
                # The converter emits malformed accent markup for some movers.
                key = "inline-" + hashlib.sha256(latex.encode()).hexdigest()[:12]
                asset = ASSETS[key]
                p.add_run().add_picture(asset["file"], width=Cm(asset["width_cm"] * size / 10.5))
                continue
            for node in wrapper:
                p._p.append(node)
        else:
            text_run(p, piece.strip("*") if piece.startswith("**") else piece,
                     size=size, bold=piece.startswith("**"))


def add_table(doc, rows):
    cols = len(rows[0])
    widths = {3: [3.0, 6.6, 7.0], 4: [5.0, 3.8, 3.9, 3.9],
              5: [2.2, 3.6, 3.6, 3.6, 3.6], 6: [3.1, 2.7, 2.7, 2.7, 2.7, 2.7]}[cols]
    if rows[0][0] == "配置":
        widths = [5.6, 5.5, 5.5]
    elif rows[0][0] == "变量":
        widths = [2.4, 5.6, 8.6]
    elif rows[0][0] in ("观测条件", "评价工况"):
        widths = [3.0, 3.4, 3.4, 3.4, 3.4]
    table = doc.add_table(rows=len(rows), cols=cols)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    for ri, values in enumerate(rows):
        common.prevent_row_split(table.rows[ri])
        for ci, value in enumerate(values):
            cell = table.cell(ri, ci)
            cell.width = Cm(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = (WD_ALIGN_PARAGRAPH.LEFT if cols == 3 and rows[0][0] != "配置"
                           else WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            inline(p, value, 8.5)
            common.set_cell_margins(cell, top=40, bottom=40, start=55, end=55)
            off = {"val": "nil"}
            common.set_cell_border(cell, top=off, bottom=off, start=off, end=off)
            if ri == 0:
                for run in p.runs:
                    run.bold = True
                common.set_cell_border(cell, top={"val": "single", "sz": "9", "color": "000000"},
                                       bottom={"val": "single", "sz": "5", "color": "000000"})
            elif ri == len(rows) - 1:
                common.set_cell_border(cell, bottom={"val": "single", "sz": "9", "color": "000000"})
    common.set_repeat_table_header(table.rows[0])


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=ROOT / "v2/paper/PHASE_PERIODIC_ENGINEERING_MANUSCRIPT_V2.md")
    parser.add_argument("--equations", type=Path, required=True)
    parser.add_argument("--output", type=Path, default=ROOT / "v2/paper/output/docx/渗透率场补全_补充实验重写稿_V2.docx")
    args = parser.parse_args()
    lines = args.source.read_text().splitlines()
    assets = json.loads((args.equations / "manifest.json").read_text())
    ASSETS.update(assets)
    doc = Document()
    style.configure(doc)
    doc.sections[0].header.paragraphs[0].clear()
    doc.styles["Normal"].paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    doc.styles["Normal"].paragraph_format.line_spacing = Pt(16)
    for name in ["Normal", "Title", "Journal Heading 1", "Journal Heading 2"]:
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    index = 0
    while index < len(lines):
        text = lines[index].strip()
        index += 1
        if not text or text.startswith("<!--"):
            continue
        if text.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            text_run(p, text[2:], 17, True)
        elif text == "$$":
            parts = []
            while lines[index].strip() != "$$":
                parts.append(lines[index])
                index += 1
            index += 1
            latex = "\n".join(parts)
            tag = re.search(r"\\tag\{(\d+)\}", latex)[1]
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.keep_together = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(assets[tag]["file"], width=Cm(min(assets[tag]["width_cm"], 15.0)))
            text_run(p, f"  ({tag})", 10)
        elif text.startswith("### "):
            common.add_heading(doc, text[4:], 2)
        elif text.startswith("## "):
            common.add_heading(doc, text[3:], 1)
        elif text.startswith("|"):
            rows, index = common.parse_table(lines, index - 1)
            add_table(doc, rows)
        elif text.startswith("!["):
            image = re.search(r"\]\(([^)]+)\)", text)[1]
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str((args.source.parent / image).resolve()), width=Cm(16.6))
        elif text.startswith("**图") or text.startswith("**表"):
            p = doc.add_paragraph(style="Journal Caption")
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.keep_with_next = text.startswith("**表")
            inline(p, text, 9)
        else:
            p = doc.add_paragraph()
            is_ref = bool(re.match(r"^\[\d+\]", text))
            if is_ref or text.startswith("**"):
                p.paragraph_format.first_line_indent = Pt(0)
            if is_ref:
                p.paragraph_format.line_spacing = Pt(13)
            inline(p, text, 9 if is_ref else 10.5)
    for element in doc.element.iter(qn("w:r")):
        text = "".join(element.itertext())
        if re.search(r"[\u3400-\u9fff]", text):
            fonts = element.get_or_add_rPr().get_or_add_rFonts()
            for key in ("ascii", "hAnsi", "eastAsia", "cs"):
                fonts.set(qn(f"w:{key}"), "Songti SC")
    doc.core_properties.title = lines[0][2:]
    doc.core_properties.author = ""
    doc.core_properties.subject = "渗透率场补全数值实验"
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
