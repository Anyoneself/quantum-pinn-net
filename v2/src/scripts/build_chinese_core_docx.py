#!/usr/bin/env python3
"""Build the Chinese journal draft using verified MathJax equation images."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import build_submission_docx as common


ROOT = Path(__file__).resolve().parents[3]
SOURCE = ROOT / "v2/paper/PHASE_PERIODIC_CHINESE_CORE_DRAFT.md"
OUTPUT = ROOT / "v2/paper/output/docx/周期相位网络_中文核心期刊风格稿.docx"


def configure(doc):
    # Retain the repository's established A4 journal layout, without its
    # previous target-journal masthead or submission-ready metadata.
    common.configure_document(doc)
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2.0)
    title = doc.styles["Title"]
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.font.size = Pt(17)
    title.paragraph_format.first_line_indent = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.keep_with_next = True
    for element in doc.styles.element.iter(qn("w:pBdr")):
        element.getparent().remove(element)
    section.header.paragraphs[0].clear()
    p = section.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    common.set_run_font(
        p.add_run("渗透率场内补全研究"), size=8,
    )
    for name in ("Journal Heading 1", "Journal Heading 2"):
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)


def equation(doc, latex, assets):
    number = re.search(r"\\tag\{(\d+)\}", latex).group(1)
    asset = assets[number]
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.keep_together = True
    width = min(asset["width_cm"], 15.0)
    p.add_run().add_picture(asset["file"], width=Cm(width))
    common.set_run_font(p.add_run(f"    ({number})"), size=10)


def table(doc, rows):
    count = len(rows[0])
    widths = {
        6: [2.4, 4.0, 2.1, 2.1, 2.4, 2.8],
        5: [2.4, 3.5, 3.5, 3.5, 3.5],
        3: [3.8, 8.0, 4.4],
    }[count]
    tab = doc.add_table(rows=len(rows), cols=count)
    tab.autofit = False
    tab.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for column, width in zip(tab.columns, widths):
        column.width = Cm(width)
    for ri, values in enumerate(rows):
        common.prevent_row_split(tab.rows[ri])
        for ci, value in enumerate(values):
            cell = tab.cell(ri, ci)
            cell.width = Cm(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            common.set_cell_margins(cell, top=45, bottom=45, start=65, end=65)
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            common.add_inline(p, value, size=9, base_bold=(ri == 0))
            # Three-line tables follow the retained Chinese journal style.
            off = {"val": "nil"}
            common.set_cell_border(
                cell, top=off, bottom=off, start=off, end=off,
                insideH=off, insideV=off,
            )
            if ri == 0:
                common.set_cell_border(
                    cell, top={"val": "single", "sz": "10", "color": "000000"},
                    bottom={"val": "single", "sz": "5", "color": "000000"},
                )
            elif ri == len(rows) - 1:
                common.set_cell_border(
                    cell, bottom={"val": "single", "sz": "10", "color": "000000"},
                )
    common.set_repeat_table_header(tab.rows[0])
    return tab


def build(source, output, equations):
    lines = source.read_text(encoding="utf-8").splitlines()
    assets = json.loads((equations / "manifest.json").read_text())
    doc = Document()
    configure(doc)
    title = lines[0].removeprefix("# ")
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    common.set_run_font(p.add_run(title), size=17, bold=True)
    index = 1
    english = False
    while index < len(lines):
        text = lines[index].strip()
        index += 1
        if not text:
            continue
        if text == r"\[":
            parts = []
            while lines[index].strip() != r"\]":
                parts.append(lines[index].strip())
                index += 1
            index += 1
            equation(doc, " ".join(parts), assets)
        elif text.startswith("|"):
            rows, index = common.parse_table(lines, index - 1)
            table(doc, rows)
        elif text.startswith("!["):
            target = re.search(r"\]\(([^)]+)\)", text).group(1)
            common.add_figure(doc, (source.parent / target).resolve())
        elif text.startswith("### "):
            common.add_heading(doc, text[4:], 2)
        elif text.startswith("## "):
            heading = text[3:]
            english = heading.startswith("Within Field")
            p = common.add_heading(doc, heading, 1)
            if english:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(14)
                for run in p.runs:
                    common.set_run_font(
                        run, size=13, bold=True, cn=common.EN_BODY,
                    )
        elif text.startswith(("**表", "**图")):
            p = doc.add_paragraph(style="Journal Caption")
            common.add_inline(p, text, size=9, base_bold=True)
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(4)
            common.set_keep_with_next(p, True)
        else:
            ref = bool(re.match(r"^\[\d+\]", text))
            note = text.startswith("注：")
            meta = text.startswith("作者、")
            size = 9 if note or ref else 10.5
            p = common.add_body_paragraph(
                doc, text, size=size,
                indent=not (note or ref or meta or english or text.startswith("**")),
            )
            p.paragraph_format.line_spacing = Pt(14 if note or ref else 16)
            if meta:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(5)
            if ref:
                p.paragraph_format.left_indent = Pt(18)
                p.paragraph_format.first_line_indent = Pt(-18)
                p.paragraph_format.space_after = Pt(2)
            if english:
                p.paragraph_format.line_spacing = Pt(14)
                for run in p.runs:
                    common.set_run_font(
                        run, size=10.5, bold=run.bold, cn=common.EN_BODY,
                    )
    doc.core_properties.title = title
    doc.core_properties.subject = "探索性场内空间补全研究"
    doc.core_properties.author = ""
    for root in (doc.element, doc.sections[0].header._element):
        for run in root.iter(qn("w:r")):
            text = "".join(run.itertext())
            family = "Songti SC" if re.search(r"[\u3400-\u9fff]", text) else "Times New Roman"
            props = run.get_or_add_rPr()
            fonts = props.get_or_add_rFonts()
            for name in ("ascii", "hAnsi", "eastAsia", "cs"):
                fonts.set(qn(f"w:{name}"), family)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--equations", type=Path, required=True)
    args = parser.parse_args()
    build(args.source.resolve(), args.output.resolve(), args.equations.resolve())


if __name__ == "__main__":
    main()
